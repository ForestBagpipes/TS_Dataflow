# -*- coding: utf-8 -*-
"""Insert paragraphs into the progress docx, matching surrounding body format."""
import copy, sys
import docx

PATH = "docx/胡宏彬-进度文档-20260818.docx"


def add_after(doc, anchor_idx, texts, template_idx=None):
    """Insert `texts` as body paragraphs right after paragraphs[anchor_idx]."""
    tmpl = doc.paragraphs[template_idx if template_idx is not None else anchor_idx]
    anchor = doc.paragraphs[anchor_idx]
    prev = anchor._p
    made = []
    for t in texts:
        new = copy.deepcopy(tmpl._p)
        # drop every run, keep paragraph properties
        for r in new.findall(docx.oxml.ns.qn('w:r')):
            new.remove(r)
        prev.addnext(new)
        prev = new
        p = docx.text.paragraph.Paragraph(new, anchor._parent)
        run_tmpl = tmpl.runs[0] if tmpl.runs else None
        r = p.add_run(t)
        if run_tmpl is not None and run_tmpl._element.rPr is not None:
            r._element.insert(0, copy.deepcopy(run_tmpl._element.rPr))
        made.append(p)
    return made


def find(doc, prefix, start=0):
    for i, p in enumerate(doc.paragraphs):
        if i >= start and p.text.strip().startswith(prefix):
            return i
    raise SystemExit(f"anchor not found: {prefix}")
