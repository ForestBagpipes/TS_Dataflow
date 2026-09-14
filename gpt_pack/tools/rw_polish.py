# -*- coding: utf-8 -*-
"""Final pass. Thins out the contrastive construction and the leftover template
phrases. A few instances are kept where the sentence genuinely corrects a
reading the reader is likely to have, which is what the construction is for."""
import sys
sys.path.insert(0, "tools")
import docx
from docx_rewrite import count_math, find_paragraph, set_paragraph_md

PATH = "docx/胡宏彬-进度文档-20260819.docx"

# 逐条替换，键为原句片段，值为改写后的整句片段
EDITS = [
    ("其能力上限越来越取决于训练语料的质量而非模型结构。",
     "模型结构趋于稳定之后，能力上限主要由训练语料的质量决定。"),
    ("这样的设定更接近上下文老虎机而非长程决策问题。",
     "这样的设定接近上下文老虎机，与长程决策问题有明显距离。"),
    ("因此本文采用这一族方法而非深度强化学习。",
     "本文因此采用这一族方法。"),
    ("三者都取自本次候选自身而非窗口的缺陷假设。",
     "三者都取自本次候选自身，与窗口的缺陷假设无关。"),
    ("把两类证据直接相加会让盲区叠加而非抵消，",
     "把两类证据直接相加，盲区会叠加起来，"),
    ("合取而非加权求和的选择可以从损害的角度得到支持。",
     "选择合取而不是加权求和，可以从损害的角度得到支持。"),
    ("策略实际收到的回报因而是经过结构条件过滤后的 ",
     "策略实际收到的回报因而是经过结构条件过滤后的 "),
    ("而非会中断决策过程的外部干预。",
     "它限制动作集合，但不中断决策过程。"),
    ("SPO 采用上下文老虎机而非深度强化学习，理由是",
     "SPO 采用上下文老虎机，理由是"),
    ("这是采用奖励均值而非接受率乘平均效用增益的理由之一，后者在无一通过时得到零，无法区分从未有效与从未尝试。",
     "这是采用奖励均值的理由之一。接受率乘平均效用增益在无一通过时得到零，从未有效与从未尝试就分不开了。"),
    ("更低的取值会压到奖励分布的主体而非尾部，改变相当一部分窗口的动作选择。",
     "更低的取值会压到奖励分布的主体，改变相当一部分窗口的动作选择。"),
    ("而不是会使决策中断的约束。",
     "它限制可选动作，决策本身不会中断。"),
    ("比较基准是屏蔽后的最优策略而非无约束的最优策略，",
     "比较基准取屏蔽后的最优策略，"),
    ("决策质量组包含四个维度，用于评估智能体本身而非其产物。",
     "决策质量组包含四个维度，评估的对象是智能体本身。"),
    ("种子写定而非运行时选取，可以避免重跑落到更有利的划分上。",
     "种子在运行之前写定，可以避免重跑落到更有利的划分上。"),
    ("若信号度量的是可预测程度而非质量，两组的读数会分居两端。",
     "若信号度量的是可预测程度，两组的读数会分居两端。"),
    ("策略学习的收益主要来自抑制无望提议而非提高修复量，",
     "策略学习的收益主要来自抑制无望提议，"),
    ("曲线的横轴取计算开销而非探测次数，",
     "曲线的横轴取计算开销，"),
    ("若屏蔽的行为随输入变化而非一律保守，",
     "若屏蔽的行为随输入变化，"),
    ("问题出在被校准的量而非校准方式，",
     "问题出在被校准的量上，"),
    ("更像分辨率不足而非结构，",
     "更像分辨率不足，"),
    ("策略学习的收益因而主要来自对插补提议的取舍，而非全部算子。",
     "策略学习的收益因而集中在插补提议的取舍上。"),
    ("这是确定性策略的结构性缺失而非数据稀疏，",
     "这是确定性策略的结构性缺失，与数据稀疏无关，"),
    ("需要说明的是 90 分位与另两档的差距略大，",
     "90 分位与另两档的差距略大，"),
    ("压缩幅度已经影响到分布主体而非仅尾部。",
     "压缩幅度已经触及分布主体。"),
    ("此外，把这套以屏蔽塑形奖励的机制推广到",
     "把这套以屏蔽塑形奖励的机制推广到"),
]


def main():
    d = docx.Document(PATH)
    before = count_math(d)
    hits = 0
    misses = []
    for old, new in EDITS:
        found = False
        for p in d.paragraphs:
            if old in p.text:
                t = p.text.replace(old, new)
                # 保留首个 run 的格式，其余清掉
                runs = p.runs
                if not runs:
                    continue
                for r in runs[1:]:
                    r._element.getparent().remove(r._element)
                runs[0].text = t
                hits += 1
                found = True
                break
        if not found:
            misses.append(old[:30])
    d.save(PATH)
    d2 = docx.Document(PATH)
    txt = "\n".join(p.text for p in d2.paragraphs)
    print(f"替换 {hits} 处，未命中 {len(misses)}")
    for m in misses:
        print("   未命中:", m)
    print(f"剩余 而非 {txt.count('而非')}，而不是 {txt.count('而不是')}，"
          f"需要说明的是 {txt.count('需要说明的是')}，此外 {txt.count('此外')}")
    print(f"公式 {before} -> {count_math(d2)}")


if __name__ == "__main__":
    main()
