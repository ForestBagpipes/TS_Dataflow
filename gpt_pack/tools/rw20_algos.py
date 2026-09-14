# -*- coding: utf-8 -*-
"""Turn the symbols inside the three algorithm listings into Word equations.

Each pseudocode line keeps its step number, its control keywords and its comment
as plain text, and every expression between them becomes an equation object.
Indentation uses a non breaking space so markdown does not swallow it and the
nesting of the loops survives.
"""
import sys
sys.path.insert(0, "tools")
import docx
from docx_rewrite import count_math, set_paragraph_md

PATH = "docx/胡宏彬-进度文档-20260820.docx"
NB = " "
I = NB * 2   # one level of indentation


def lines_of(table):
    return [p for r in table.rows for c in r.cells
            for p in c.paragraphs if p.text.strip()]


ALG1 = {
 "输入：": r"输入：窗口 $x$，冻结模型池 $\mathcal{F}$，邻域大小 $K$",
 "输出：": r"输出：状态 $s$，画像簇标识 $j$",
 "1:": r"1: $p\leftarrow\mathrm{ExtractProfile}(x)$" + NB * 6 + "// 十二维统计画像",
 "2:": r"2: $b\leftarrow\mathrm{Probe}(f_{1},x)$" + NB * 8 + "// 前向探测得到行为签名",
 "3:": r"3: for $f\in\mathcal{F}\setminus\{f_{1}\}$ do "
       r"$b\leftarrow b\cup\mathrm{Disagree}(f,x)$ end for" + NB * 2 + "// 模型间分歧",
 "4:": r"4: $\mathcal{N},j\leftarrow\mathrm{KNN}(p,K)$" + NB * 5 + "// 同类邻居与所属画像簇",
 "5:": r"5: $z\leftarrow\mathrm{RobustZ}(b,\mathcal{N})$" + NB * 5
       + "// 以邻居的中位数与四分位距标准化",
 "6:": r"6: $r_{\mathrm{stat}}\leftarrow\mathrm{ScoreStat}(p)$，"
       r"$r_{\mathrm{behav}}\leftarrow z^{\top}w$",
 "7:": r"7: $h\leftarrow\mathrm{Posterior}(r_{\mathrm{stat}},r_{\mathrm{behav}})$"
       + NB * 3 + "// 缺陷类型后验",
 "8:": r"8: return $s\leftarrow(p,b,r_{\mathrm{stat}},r_{\mathrm{behav}},h,\varnothing,B)$，$j$",
}

ALG2 = {
 "输入：": r"输入：工作副本 $x'$，候选动作 $a$，阈值 $\varepsilon$、$\tau$、$\eta$",
 "输出：": r"输出：裁决结果、更新后的副本与奖励 $r$",
 "1:": r"1: $c\leftarrow a(x')$" + NB * 8 + "// 作用于副本的拷贝",
 "2:": r"2: $\Delta U\leftarrow U(c)-U(x')$" + NB * 5 + "// 同跨度同尺度重测",
 "3:": r"3: $d\leftarrow D(x',c)$，$\rho\leftarrow R(a)$",
 "4:": r"4: if $\Delta U>\varepsilon$ and $d<\tau$ and $\rho<\eta$ then",
 "5:": "5:" + I + r"$x'\leftarrow c$，$\mathrm{verdict}\leftarrow\mathrm{ACCEPT}$，"
       r"$r\leftarrow\Delta U$",
 "6:": "6: else",
 "7:": "7:" + I + r"$\mathrm{verdict}\leftarrow\mathrm{ROLLBACK}$，"
       r"$r\leftarrow-c_{0}\mathrm{Cost}(a)$" + NB * 2 + "// 副本不变",
 "8:": "8:" + I + r"$\mathrm{reason}\leftarrow$ 首个不满足的条件" + NB * 2 + "// 供策略调整方向",
 "9:": "9: end if",
 "10:": r"10: return $\mathrm{verdict}$，$x'$，$r$，$\mathrm{reason}$",
}

ALG3 = {
 "输入：": r"输入：语料 $\mathcal{D}$，总预算 $B$，探索系数 $c_{u}$，"
          r"标定间隔 $T_{\mathrm{cal}}$，注入概率 $p_{\mathrm{inj}}$",
 "输出：": r"输出：治理后语料 $\mathcal{D}'$，价值表 $Q$",
 "1:": r"1: 由固定规则策略的历史轨迹暖启动 $Q$ 与 $n$，未访问格取乐观初值，$t\leftarrow0$",
 "2:": r"2: for $x\in\mathcal{D}$ do",
 "3:": "3:" + I + r"$s,j\leftarrow\mathrm{CDP}(x)$" + NB * 4 + "// 状态与画像簇",
 "4:": "4:" + I + r"$b\leftarrow\mathrm{AllocBudget}(Q,j,s,B)$" + NB * 2
       + "// 按期望收益分配预算",
 "5:": "5:" + I + r"while $b>0$ do",
 "6:": "6:" + I * 2 + r"$\mathcal{A}_{\mathrm{shield}}\leftarrow\mathrm{Shield}(s)$"
       + NB * 2 + "// 由定理 2 保证非空",
 "7:": "7:" + I * 2 + r"$\mathcal{C}\leftarrow\mathrm{Inject}"
       r"(\mathcal{A}_{\mathrm{shield}},j,p_{\mathrm{inj}})$" + NB * 2
       + "// 未访问格按概率强制入选",
 "8:": "8:" + I * 2 + r"$a\leftarrow\arg\max_{a\in\mathcal{C}}"
       r"[Q(j,a)+c_{u}\sqrt{2\ln t/n(j,a)}]$",
 "9:": "9:" + I * 2 + r"if $a\in\{\mathrm{KEEP},\mathrm{ABSTAIN}\}$ then break",
 "10:": "10:" + I * 2 + r"$\mathrm{verdict},x',r,\mathrm{reason}\leftarrow\mathrm{CDS}(x',a)$"
        + NB * 2 + "// 沙箱执行与裁决",
 "11:": "11:" + I * 2 + r"$n(j,a)\leftarrow n(j,a)+1$，$t\leftarrow t+1$",
 "12:": "12:" + I * 2 + r"$Q(j,a)\leftarrow Q(j,a)+[r-Q(j,a)]/n(j,a)$" + NB * 2
        + "// 增量更新",
 "13:": "13:" + I * 2 + r"$s\leftarrow\mathrm{Update}(s,\mathrm{verdict},\mathrm{reason})$，"
        r"$b\leftarrow b-\mathrm{Cost}(a)$",
 "14:": "14:" + I + r"end while",
 "15:": "15:" + I + r"if $t\bmod T_{\mathrm{cal}}=0$ then "
        r"$\tau\leftarrow\mathrm{Calibrate}(\alpha)$ end if",
 "16:": "16: end for",
 "17:": r"17: return $\mathcal{D}'$，$Q$",
}


def apply(table, mapping, label):
    ps = lines_of(table)
    done, missed = 0, []
    used = set()
    for p in ps:
        head = p.text.strip()
        key = None
        for k in mapping:
            if head.startswith(k) and k not in used:
                if key is None or len(k) > len(key):
                    key = k
        if key is None:
            continue
        set_paragraph_md(p, mapping[key])
        used.add(key)
        done += 1
    missed = [k for k in mapping if k not in used]
    print(f"{label}: 改写 {done} 行" + (f"，未匹配 {missed}" if missed else ""))
    return done


def main():
    d = docx.Document(PATH)
    before = count_math(d)
    apply(d.tables[3], ALG1, "算法 1")
    apply(d.tables[4], ALG2, "算法 2")
    apply(d.tables[5], ALG3, "算法 3")
    d.save(PATH)
    d2 = docx.Document(PATH)
    print(f"公式 {before} -> {count_math(d2)}")


if __name__ == "__main__":
    main()
