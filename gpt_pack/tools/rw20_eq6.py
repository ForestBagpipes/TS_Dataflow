# -*- coding: utf-8 -*-
"""Complete the formulas that named a quantity without defining it.

Equation 6, the decision risk, weighted three terms whose own definitions
appeared nowhere. Depth and consistency are now written out, matching
verify.improvement_depth and verify.improvement_consistency. The same pass
gives the structural difference per attribute a definition and states what the
cost estimate in the budget rule actually is.
"""
import sys
sys.path.insert(0, "tools")
import docx
from docx_rewrite import (count_math, find_paragraph, insert_after,
                          set_paragraph_md)

PATH = "docx/胡宏彬-进度文档-20260820.docx"


def main():
    d = docx.Document(PATH)
    before = count_math(d)

    def rw(prefix, md):
        set_paragraph_md(d.paragraphs[find_paragraph(d, prefix)], md)

    # 公式 6 决策风险，补齐两个分量的定义
    i = find_paragraph(d, "其中  为决策风险，度量当前证据支持该动作的充分程度")
    set_paragraph_md(d.paragraphs[i],
        r"其中 $R(a)$ 为决策风险，取值在 $[0,1]$，越小表示提交越安全，由三项加权而成")
    j = find_paragraph(d, "式中")
    set_paragraph_md(d.paragraphs[j],
        r"式中 $\mathrm{cost}(a)$ 是算子的归一化代价，"
        r"$\mathrm{depth}$ 与 $\mathrm{cons}$ 由校准后行为分量在干预前后的取值给出。"
        r"记干预前后的标准化行为向量为 $z$ 与 $z'$，各分量都是越小越好，"
        r"则某一维的改善量为 $\delta_{k}=z_{k}-z'_{k}$。一致性是改善方向正确的维度占发生变化维度的比例")
    insert_after(d, j, [
        r"$$\mathrm{cons}=\frac{|\{k:\delta_{k}>0\}|}"
        r"{\max(|\{k:\delta_{k}\neq0\}|,1)}$$",
        r"深度则取改善维度上改善量的中位数，并截断到 $[0,1]$",
        r"$$\mathrm{depth}=\mathrm{clip}\big(\mathrm{med}\{\delta_{k}:\delta_{k}>0\},0,1\big)$$",
        "两者分别回答改善铺开了多少维度与这些维度改善了多少。"
        "只看铺开的宽度会把普遍而微弱的变化误判为有效，"
        "只看幅度则会让单一维度的塌陷主导判断，因此两者同时进入风险。"
        "一个标准差单位的中位改善已经是决定性的编辑，深度在该处饱和。"
        "三项都取自本次候选自身，与窗口的缺陷假设无关。"
        "用缺陷假设的后验充当风险曾是早期做法，实测显示它与编辑是否成功反向，"
        "后验最确信的十分之一窗口反而更容易收到错误编辑，因此改用上面的三项。"
        "上式定义在单个窗口上，而治理的最终对象是整个语料，"
        "两者通过下面的损害定义连接。",
    ], template_idx=j)

    # 公式 3 的属性差异，给出定义
    rw("加权平均之外保留最差单项",
       r"式中 $d_{j}$ 是第 $j$ 个属性在编辑前后的取值经该属性自身量纲归一化后的绝对差，"
       r"取值在 $[0,1]$，权重 $w_{j}$ 按属性对该类算子的相关性给定。"
       "加权平均之外保留最差单项，可以避免某一项上的灾难性破坏被其余项平均掉。"
       "序列尾部被整体抹平就属于这种情形，它在多数属性上看起来变化不大，只在极值一项上暴露。"
       r"$D$ 与模型无关，因而能够独立于效用信号提供约束，在策略学习过程中也不会被策略操纵。")

    # 公式 11 的代价估计说明
    rw("缺陷类型后验接近确定且指向干净的窗口获得零预算",
       r"式中 $B_{\mathrm{left}}$ 为尚未分配的预算，"
       r"$\hat{c}_{i}$ 为窗口 $i$ 的预估探测代价，当前实现取语料的平均探测次数作为该估计，"
       "在同一簇内积累足够观测之后可以改用该簇的历史均值。"
       "缺陷类型后验接近确定且指向干净的窗口获得零预算并被直接跳过，"
       "后验分散的窗口则获得更多次尝试机会。"
       "这一机制把探测集中在证据模糊的地方，其效果在 4.2 节以预算效率曲线衡量。"
       "算法 3 给出完整流程。")

    d.save(PATH)
    d2 = docx.Document(PATH)
    print(f"公式补齐完成，公式 {before} -> {count_math(d2)}")


if __name__ == "__main__":
    main()
