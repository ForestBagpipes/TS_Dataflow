# -*- coding: utf-8 -*-
"""Rebuild the seven paragraphs the plain text pass flattened.

The polish pass wrote its replacement into the first run and deleted the others,
which is safe for a paragraph of pure text and wrong for one holding equations,
because an oMath element is not a run and therefore survived in place while the
words around it collapsed to the front. These seven are rebuilt from markdown so
the equations land where the sentence puts them.
"""
import sys
sys.path.insert(0, "tools")
import docx
from docx_rewrite import count_math, find_paragraph, set_paragraph_md

PATH = "docx/胡宏彬-进度文档-20260819.docx"

REBUILD = [
    ("式中",
     r"式中 $\mathrm{depth}$ 为行为证据改善的深度，$\mathrm{cons}$ 为改善在各行为分量上的一致性，"
     r"$\mathrm{cost}$ 为算子的归一化代价，三者都取自本次候选自身，与窗口的缺陷假设无关。"
     r"用假设的后验充当风险曾是早期做法，实测显示它与编辑是否成功反向，"
     r"后验最确信的十分之一窗口反而更容易收到错误编辑，因此改用上式的三项。"
     r"$R(a)$ 越小表示提交越安全。上式定义在单个窗口上，"
     "而治理的最终对象是整个语料，两者通过下面的损害定义连接。"),

    ("把两类证据直接相加",
     "把两类证据直接相加，盲区会叠加起来，因为它们的量纲和分布都不同。"
     r"CDP 的处理是先在画像空间中检索与当前窗口最相似的 $K$ 个邻居 $\mathcal{N}(x)$，"
     "相似度用画像向量经 L2 归一化后的余弦距离衡量，"
     "再用这组邻居的中位数与四分位距对行为签名做稳健标准化"),

    ("屏蔽的第二重作用来自它与奖励的耦合",
     "屏蔽的第二重作用来自它与奖励的耦合。由 3.2 节的奖励定义，"
     "只有通过屏蔽的动作才计入正奖励，被撤销的动作只留下探测代价。"
     r"策略实际收到的回报因而是经过结构条件过滤后的 $\Delta U$，不是原始读数。"
     "一个抹平动作即使拿到很高的效用读数，也会因结构失真超限而被撤销，最终计入负奖励。"
     "策略在若干轮之后会学到，在需要保护的窗口上提出抹平动作是不划算的，"
     r"从而在提议阶段就避开这类动作。若把奖励换成未经过滤的 $\Delta U$，"
     "同样的学习过程会得到相反的结果。这一对比构成了 4.2 节奖励塑形消融实验的设计依据。"),

    ("接受奖励需要设上界",
     "接受奖励需要设上界。定理 5 的遗憾界要求奖励有界或次高斯，"
     "而实测的效用变化两者皆非，其分布呈长尾，"
     "少数窗口的读数比中位数高出两个数量级，远超上置信界在同等规模下的探索宽度。"
     "若不设上界，持有极端值的格子将被永久选中而探索停止，这破坏的是定理的前提。"
     r"因此接受奖励按训练划分上的分位数 $q$ 做单侧截断，即 $r=\min(\Delta U,q)$，"
     "评测划分不参与该数值的确定。撤销奖励本身已由代价系数限定，不需要截断。"
     "分位数取值经过敏感性检验，取值不宜低于 95 分位，"
     "更低的取值会压到奖励分布的主体，改变相当一部分窗口的动作选择。"),

    ("证明。由定理 2",
     r"证明。由定理 2，$\mathcal{A}_{\mathrm{shield}}(s)$ 在任意状态下非空，"
     "因此屏蔽构成一个状态相关的动作集限制，它限制可选动作，决策本身不会中断。"
     "带受限动作集的上下文老虎机的上置信界分析在此直接适用[19]，"
     "其遗憾界形式与动作集大小无关，只依赖特征维度与轮数。"
     "将比较基准取为屏蔽后可行动作中的最优策略，即得结论。证毕。"),

    ("定理 5 说明屏蔽在提供安全保证的同时",
     "定理 5 说明屏蔽在提供安全保证的同时保留了策略学习的渐近性质。"
     "比较基准取屏蔽后的最优策略，它与无约束最优策略之间的差距对应屏蔽为安全付出的代价，"
     r"其大小由 $\tau$ 决定，而 $\tau$ 又由目标损害率 $\alpha$ 确定。"
     "这一设计把安全与效率的权衡集中在一个具有明确统计含义的参数上。"),

    ("这一组用来确认问题设定成立",
     "这一组用来确认问题设定成立，即模型行为信号在部分干净数据上会给出与有害性相反的读数。"
     "干净分布外层由随机游走、阶梯、脉冲串与锯齿四种合成形态构成，"
     "真实跨域层取自 Exchange 且不做注入，两组数据都是干净的且都属于分布外。"
     "按形态分别统计行为风险的分布并作对照，若信号度量的是可预测程度，"
     "两组的读数会分居两端。同时统计各层的编辑经济学，"
     r"即每破坏单位方差换到的效用增益 $\Delta U/\Delta\mathrm{Var}$。"),
]


def main():
    d = docx.Document(PATH)
    before = count_math(d)
    for prefix, md in REBUILD:
        set_paragraph_md(d.paragraphs[find_paragraph(d, prefix)], md)
    d.save(PATH)
    d2 = docx.Document(PATH)
    print(f"重建 {len(REBUILD)} 段，公式 {before} -> {count_math(d2)}")


if __name__ == "__main__":
    main()
