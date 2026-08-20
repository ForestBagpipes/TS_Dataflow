# -*- coding: utf-8 -*-
"""Remove the method comparison table and the hyperparameter table.

Both were lists of attributes with one row per item. Their content moves into
prose, the baselines as numbered sentences and the parameters as a paragraph
that says what each one does and how to set it. The dataset table stays, it is
the one table the experiment chapter keeps.
"""
import sys
sys.path.insert(0, "tools")
import docx
from docx_rewrite import (count_math, delete_paragraph, find_paragraph,
                          insert_after, set_paragraph_md)

PATH = "docx/胡宏彬-进度文档-20260820.docx"


def drop_table(doc, idx):
    tbl = doc.tables[idx]._tbl
    tbl.getparent().remove(tbl)


def main():
    d = docx.Document(PATH)
    before = count_math(d)
    kill = []

    def rw(prefix, md):
        set_paragraph_md(d.paragraphs[find_paragraph(d, prefix)], md)

    # 一 第 1 章的方法对比表改为正文
    rw("把治理做成自动化并不等于把清洗算子串起来",
       "把治理做成自动化并不等于把清洗算子串起来。"
       "趋势转折、状态切换、极端峰值与分布外样本同样会让预测误差升高，"
       "仅凭误差无法把它们与真正的污染分开。"
       "一个只追求降低误差的流程会优先修改这些片段，而它们恰恰是应当保留的。"
       "本文因此关心如何区分有害的污染与困难但正确的数据，并只对确实需要处理的部分动手。")

    i = find_paragraph(d, "表 1 中的方法沿两条路线展开")
    set_paragraph_md(d.paragraphs[i],
       "已有工作沿两条路线展开，按产出形式与是否具备执行后的核对能力可以逐一说明。")
    insert_after(d, i, [
        "（1）SCREEN[5]假设相邻时间步的取值变化存在上下界，"
        "把违反约束的点修复到最近的合法取值，产出修复后的数据，算子由固定规则选定。"
        "（2）IMR[6]用自回归模型逐点估计并迭代收敛，同样产出修复数据，"
        "需要少量已知正确的点作为起始约束。"
        "（3）SpeedAcc[7]在速度之外引入加速度约束，用二阶信息约束修复的平滑程度。"
        "（4）MTCSC[8]把速度约束与聚类结合，利用维度之间的相关性提高修复精度。"
        "以上四项都在算子选定后直接覆盖原值，此后没有机制核对这次改写是否达到预期，"
        "也不涉及探索，因而不存在探索是否有损的问题。",
        "（5）Data Shapley[9]用合作博弈中的 Shapley 值刻画样本价值，产出样本分数。"
        "（6）Data-OOB[10]借助袋装模型的袋外估计把开销降到可接受范围，同样产出分数。"
        "（7）TimeInf[11]把影响函数定义在时间块上再聚合到时间点，"
        "在给出逐点分数的同时保留时序结构。"
        "（8）TSRating[12]从四个维度获取大模型的成对比较，"
        "再用元学习把这些判断蒸馏成轻量评分器。"
        "这四项不产生动作，只给排序，因此无从核对某次处理是否奏效。",
        "（9）近期出现的清洗智能体[13]用分层策略编排清洗算子，"
        "策略可以从下游反馈中改进，是这两条路线之外的第三种做法。"
        "它的算子在训练过程中立即作用于原数据，探索阶段的错误无法撤回，"
        "安全性只能体现在训练收敛后的期望意义上。"
        "（10）本文的 IntroAct-TS 同样产出修复数据并由屏蔽强化学习给出策略，"
        "区别在于每次干预之后都重新测量并据此裁决，试错发生在副本上因而探索无损，"
        "同时在语料内分配探测预算。"
        "这三项能力是前九项都不具备的，也是本文要补上的缺口。",
    ], template_idx=i)
    kill.append("表 1 面向时序数据的治理与选择方法对比")

    # 二 4.1.5 超参数表改为正文
    rw("方法涉及的超参数见表 3",
       "方法涉及的超参数可以按它们约束什么来说明。"
       r"同类校准的邻域大小 $K$ 决定行为签名与多少个结构相近的窗口比较，初值取 50，"
       "语料同质性高时可以减小，异质性高时应当增大。"
       r"效用改善的最小幅度 $\varepsilon$ 是接受一次编辑所需的最低效用增益，"
       "初值按代理模型上同一批窗口的效用读数标准差的十分之一估计，"
       "这样既能滤掉噪声也不至于挡住真实改善。"
       r"决策风险上界 $\eta$ 初值取 0.5，受保护层误编辑偏多时调低。"
       r"目标损害率 $\alpha$ 初值取 0.02，按下游任务对损害的容忍度设定，"
       r"结构失真阈值 $\tau$ 由它经共形标定确定，不需要人工调节。"
       r"上置信界的探索系数 $c_{u}$ 初值取 1.0，语料规模大时可以调低以加快收敛。"
       r"探测代价的折算系数 $c_{0}$ 初值取 0.01，探测开销紧张时调高。"
       r"重新标定的间隔 $T_{\mathrm{cal}}$ 初值取每五百次决策，策略变化快时应当缩短。"
       r"模型池大小 $M$ 初值取 3，增大可以提高分歧信号的质量，代价是开销线性增长。")
    kill.append("表 3 超参数与调节建议")

    for prefix in kill:
        try:
            delete_paragraph(d.paragraphs[find_paragraph(d, prefix)])
        except LookupError:
            print("  未找到表题:", prefix)

    # 先删超参数表再删方法表, 索引从大到小
    drop_table(d, 7)
    drop_table(d, 0)

    d.save(PATH)
    d2 = docx.Document(PATH)
    print(f"表格 {len(d.tables) + 2} -> {len(d2.tables)}，公式 {before} -> {count_math(d2)}")
    for i, tb in enumerate(d2.tables):
        print(f"  表{i} {len(tb.rows)}x{len(tb.columns)} "
              f"{tb.rows[0].cells[0].text.strip()[:30]}")


if __name__ == "__main__":
    main()
