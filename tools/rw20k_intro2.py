# -*- coding: utf-8 -*-
"""Limitations as enumerated points, then a challenge paragraph, then the method.

The reference paper states a heading sentence, enumerates two limitations each
with named counterexamples, then spends a whole paragraph on why a method that
fixes them is hard to build, and only then introduces its own. That ordering
matters because the reader needs to know the difficulty before being shown the
design, otherwise the design reads as arbitrary.
"""
import sys
sys.path.insert(0, "tools")
import docx
from docx_rewrite import (count_math, delete_paragraph, find_paragraph,
                          insert_after, set_paragraph_md)

PATH = "docx/胡宏彬-进度文档-20260820.docx"


def main():
    d = docx.Document(PATH)
    before = count_math(d)

    def rw(prefix, md):
        set_paragraph_md(d.paragraphs[find_paragraph(d, prefix)], md)

    # 局限总起
    rw("三条路线各有代价，而代价的形状恰好互补",
       "经考察，这些方法在各自设定下都能取得不错的效果，"
       "但面对本文关心的问题，普遍存在以下两方面局限。")

    i = find_paragraph(d, "经考察，这些方法在各自设定下都能取得不错的效果")
    insert_after(d, i, [
        "1）判定时刻早于执行时刻，选错之后无法挽回。"
        "修复式清洗在规则给出的那一刻就把改写作用到原值上，"
        "此后没有任何机制核对这次改写是否达到预期，"
        "SCREEN[5]与 MTCSC[8]都是如此，速度约束一旦判定某点越界，替换随即发生。"
        "学习式编排虽然让策略可以从反馈中改进，"
        "但清洗智能体[13]的算子同样在训练期间立即生效，"
        "每一次错误探索都真实地落在语料上，策略学得越多语料损坏越多。"
        "打分式选择避开了这个问题，代价是它根本不动手，"
        "面对一段只在局部含有尖峰的窗口，"
        "Data Shapley[9]与 TimeInf[11]这类方法只能整体丢弃，"
        "尖峰之外完好的部分一并舍去。"
        "三条路线因此陷入同一个两难，动手的没有回头路，有回头路的不动手。",

        "2）判定所依据的信息单一，而在需要保护的数据上这一信息会给出相反的读数。"
        "修复式清洗只看取值是否满足约束，"
        "而一次真实的极端事件在取值上与故障同样越界，约束无法把两者分开。"
        "打分式选择与学习式编排都以模型表现作为最终依据，"
        "问题在于让模型的预测变准存在两条途径，"
        "修好缺陷是一条，把序列抹平使它更容易预测是另一条，"
        "后者同样会让误差下降。"
        "在稀有事件所在的窗口上，这两条途径给出的读数方向相反，"
        "单凭模型的意见去判定，恰恰会放行那些最该拦下的修改。",
    ], template_idx=i)

    # 挑战段
    j = find_paragraph(d, "2）判定所依据的信息单一")
    insert_after(d, j, [
        "针对上述局限，当前需要一种能够在动作执行之后判定、并且判错可以撤回的治理方法，"
        "而这样一种方法的构筑面临多重挑战。"
        "一方面，判定依据的来源本身存在偏差。"
        "修改是否有害只能从效果读出，而能报告效果的仪器只有目标模型，"
        "上述第二条局限说明这台仪器在需要保护的数据上并不可信，"
        "因此判定不能只问模型，还需要一个不依赖模型的证据与之相互印证，"
        "而这个证据的阈值又无法由领域知识事先给定。"
        "另一方面，试错的代价落在不可恢复的对象上。"
        "让策略从判定结果中学习需要反复尝试，"
        "而语料与游戏或仿真环境不同，它没有可以随意重来的副本，被抹平的峰值无法恢复。"
        "如何让学习过程既能积累足够的经验，又不在积累的过程中损坏语料，"
        "是把判定推迟到执行之后所必须一并解决的问题。",
    ], template_idx=j)

    # 原来的决定与两问题合并段删除，其内容已进入挑战段
    delete_paragraph(d.paragraphs[find_paragraph(d, "把判定移到执行之后")])

    # 方法段扩写，按 PROCore 的先总后分
    rw("按照这一思路，本文提出屏蔽强化学习驱动的时序数据治理智能体",
       "为解决判定依据存在偏差与试错代价不可恢复这两个问题，"
       "本文提出屏蔽强化学习驱动的时序数据治理智能体 IntroAct-TS，"
       "把治理建模为带预算的受约束序贯决策。"
       "IntroAct-TS 首先从窗口自身的统计画像与冻结基础模型的行为反应中取证，"
       "并以画像相近的邻居为参照做稳健标准化，使不同窗口之间的读数可比。"
       "随后把候选动作施加在工作副本的拷贝上试执行并重新测量，原值在此期间保持不变，"
       "这样每一次尝试都有确切的后果可读，而尝试本身不留下痕迹。"
       "针对判定依据存在偏差这一挑战，"
       "IntroAct-TS 要求模型效用改善与时序结构受控两个相互独立的条件同时成立才允许提交，"
       "其中结构条件完全不看模型意见，"
       "并用共形风险控制把它的阈值标定到给定的损害率上，"
       "使阈值的设定不再依赖领域知识。"
       "针对试错代价这一挑战，"
       "由于所有候选都在副本上执行且提交需经判定，"
       "未通过的尝试不改变任何数据，学习过程因而不以破坏语料为代价。"
       "最后，每一次判定的结果回流为策略的奖励，"
       "被撤销的动作只留下探测代价而拿不到正回报，"
       "策略据此更新动作价值估计并在窗口之间分配探测预算，"
       "从而提高了治理在结构复杂且含噪的时序语料上的可靠性与效率。")

    d.save(PATH)
    d2 = docx.Document(PATH)
    print(f"局限挑战与方法段完成，公式 {before} -> {count_math(d2)}")


if __name__ == "__main__":
    main()
