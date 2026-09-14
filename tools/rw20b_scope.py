# -*- coding: utf-8 -*-
"""Simplify the corruption taxonomy and swap in the two scenario corpora.

The old design injected seven defect types and built six evaluation strata, four
of them synthetic. Real corpora already carry rare and difficult windows, so
constructing them is both redundant and open to the charge that the paper
designs both the fault and the thing it protects. What remains is noise
injection in two forms, random on a single channel and systematic across
channels, which is the split the reference paper uses and which maps onto how
the two failure modes actually arise.
"""
import sys
sys.path.insert(0, "tools")
import docx
from docx_rewrite import (count_math, delete_paragraph, find_paragraph,
                          insert_after, set_paragraph_md)

PATH = "docx/胡宏彬-进度文档-20260820.docx"


def main():
    d = docx.Document(PATH)
    before = count_math(d)

    def rw(prefix, md):
        set_paragraph_md(d.paragraphs[find_paragraph(d, prefix)], md)

    # 2.2 缺陷说明改为两类噪声
    rw("本文处理时序特征层面的缺陷，按影响范围分为两类",
       "本文处理时序特征层面的噪声，按其产生机制分为随机与系统两类，"
       "这一划分决定了它们在多通道语料中留下的痕迹是否可以相互印证。"
       "随机噪声来自环境扰动、设备误差与记录不一致，作用在单个通道的有限时间范围上，"
       "表现为孤立尖峰、散点缺失、连续缺失片段与重复片段。"
       "这类噪声在通道之间彼此独立，因此一个通道上的异常在同一时刻的其余通道上找不到对应。"
       "系统噪声来自采集设备的系统性偏差、数据来源切换与传感器漂移，"
       "它同时作用在多个通道上并改变整段窗口的分布，"
       "表现为水平位移、噪声放大与平台化。"
       "这类噪声在通道之间高度相关，同一时刻的多个通道会一起偏移。")
    rw("这两类缺陷带来的困难体现在两个层面",
       "两类噪声给治理带来的困难并不相同。"
       "随机噪声在单个窗口上留下明显的局部痕迹，统计画像对其检测能力较强，"
       "困难在于它与真实发生的极端事件在波形上几乎一致，"
       "仅凭单通道证据无法判断某个尖峰是故障还是事件。"
       "系统噪声改变整段分布，在模型行为上的表现与合法的状态切换高度相似，"
       "水平位移与设备检修后的基线变化都会抬高预测误差，仅凭模型反馈同样无法区分。"
       "在语料的尺度上还有一重困难，模型行为信号反映的是窗口对该模型的可预测程度，"
       "而窗口的可预测程度与它的质量并不等同，"
       "结构上本就复杂却完全干净的数据同样会得到高风险读数。"
       "这一偏差对基于规则的方法只是误判，对基于学习的方法则更严重，"
       "因为策略会主动朝着偏差的方向优化，把降低可预测难度当作目标去追求。")

    # 4.1.2 注入与分层
    rw("无标签场景下无法直接获得治理是否正确的判据",
       "无标签场景下无法直接获得治理是否正确的判据，"
       "因此在干净窗口上注入受控噪声并逐条记录标签，把注入位置作为治理应当处理的对象。"
       "注入按 2.2 节的划分分为两类。"
       "随机噪声在单个通道上注入，位置与幅度独立随机，"
       "覆盖孤立尖峰、散点缺失、连续缺失片段与重复片段四种形态。"
       "系统噪声在同一时间位置的多个通道上同时注入，"
       "覆盖水平位移、噪声放大与平台化三种形态，同一次注入的多个通道共享偏移方向。"
       "注入比例设置从百分之五到百分之六十七的多个档位，采用嵌套构造，"
       "即高比例语料包含低比例语料的全部受污染窗口，干净部分保持不变，"
       "使不同比例之间的差异只来自比例本身。")
    rw("评测集分为六层",
       "评测集按窗口的来源与是否被注入分为四层。"
       "注入层是施加了上述两类噪声的窗口，治理应当处理它们。"
       "干净层未做任何注入，治理不应改动它们。"
       "稀有层从真实数据中筛选，取的是取值落在分布尾部但完全未经注入的窗口，"
       "金融语料中的剧烈波动时段与工业语料中的启停时段都属于此类。"
       "困难层同样来自真实数据，取季节强度与自相关都较低因而难以预测但完全干净的窗口。"
       "后两层与干净层统称受保护层，其上的任何编辑都计为损害。"
       "稀有层与困难层取自数据本身而不是合成构造，"
       "这样保护对象的定义不依赖于注入过程，避免了设计缺陷同时又设计保护对象所带来的循环。")

    # 4.1.1 数据集
    rw("计划使用的数据集见表 1",
       "计划使用的数据集见表 1，覆盖工业与金融两类场景。"
       "选择两类而不是一类，是因为本文的同类校准与弃权机制都以语料存在结构异质性为前提，"
       "单一来源的语料无法检验这一前提是否被满足，这一点在 4.2 节的对应实验中展开。"
       "工业场景取电力变压器监测数据 ETTh1、ETTh2 与 ETTm1，"
       "它们是长期时序预测的通用基准[23]，采样频率为小时与十五分钟，"
       "多个通道来自同一台设备的不同测点，因此系统噪声在其上有明确的物理对应。"
       "金融场景取 TIME 基准[28]中的三项，"
       "Crypto 是四种加密资产的日频价格，US Term Structure 是美国国债期限结构的四十维工作日序列，"
       "Oil Price 是原油与成品油的十二维工作日价格。"
       "三者的通道之间都存在真实的联动关系，"
       "价格的共同波动与利率曲线的整体平移都是系统性的，"
       "与工业语料中的设备漂移在统计形态上相近而成因完全不同。"
       "两类场景的窗口在趋势强度、季节性与波动聚集上差异明显，"
       "合并为一个语料之后异质性显著高于任何单一来源。")

    d.save(PATH)
    d2 = docx.Document(PATH)
    print(f"缺陷口径与数据集改写完成，公式 {before} -> {count_math(d2)}")


if __name__ == "__main__":
    main()
